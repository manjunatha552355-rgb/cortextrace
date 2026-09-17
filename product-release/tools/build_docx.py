"""Builds the editable DOCX from build/doc-model.json (same content model as the PDF)."""
import json, re, pathlib
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parent.parent
model = json.loads((ROOT / "build" / "doc-model.json").read_text(encoding="utf-8"))
meta = model["meta"]
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x1F, 0x5F, 0xAF)
GREY = RGBColor(0x4B, 0x55, 0x63)
FONT = "Segoe UI"
MONO = "Consolas"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
TEXT_WIDTH = sec.page_width - sec.left_margin - sec.right_margin

def set_font(style, size, bold=False, color=None):
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

styles = doc.styles
set_font(styles["Normal"], 10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for name, size, before in (("Heading 1", 20, 0), ("Heading 2", 13, 14), ("Heading 3", 11, 10)):
    set_font(styles[name], size, True, NAVY)
    styles[name].paragraph_format.space_before = Pt(before)
    styles[name].paragraph_format.space_after = Pt(6)
    styles[name].paragraph_format.keep_with_next = True
set_font(styles["Caption"], 8.5, False, GREY)
styles["Caption"].font.italic = False
for lst in ("List Bullet", "List Number"):
    set_font(styles[lst], 10)

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def cell_borders(table, color="D5DCE5"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if edge in ("top", "bottom", "insideH") else "nil")
        el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)

TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+?\]\((?:https?://|mailto:)[^)]+\))")

def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = MONO; r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
            r.font.size = Pt((size or 10) - 1); r.font.color.rgb = NAVY
            continue
        elif TOKEN.fullmatch(part) and part.startswith("["):
            m = re.match(r"\[(.+?)\]\((.+?)\)", part)
            rid = par.part.relate_to(m.group(2), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
            hr = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
            col = OxmlElement("w:color"); col.set(qn("w:val"), "1F5FAF"); rPr.append(col)
            if size:
                sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
            hr.append(rPr); t = OxmlElement("w:t"); t.text = m.group(1); t.set(qn("xml:space"), "preserve"); hr.append(t); h.append(hr)
            par._p.append(h)
            continue
        else:
            r = par.add_run(part)
        if size: r.font.size = Pt(size)
        if color is not None and not part.startswith("**"): r.font.color.rgb = color
        if bold: r.bold = True

def para(text, style=None, size=None, color=None, align=None, after=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size, color)
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def field(par, instr):
    r = par.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; r._r.append(el)
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = text; r._r.append(el)
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); r._r.append(el)
    return r

def table(head, rows, widths=None, font=8.8):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_borders(t)
    total = TEXT_WIDTH
    ws = [int(total * (w / 100)) for w in widths] if widths else [int(total / len(head))] * len(head)
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.width = ws[i]; shade(c, "0B1F3A")
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        add_runs(p, h, font, RGBColor(0xFF, 0xFF, 0xFF), True)
    trPr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = ws[i]
            if ri % 2: shade(cells[i], "F7F9FB")
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_runs(p, val, font)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def note(title, text, kind):
    t = doc.add_table(rows=1, cols=1); t.autofit = False
    c = t.rows[0].cells[0]; c.width = TEXT_WIDTH
    shade(c, "FDF6EC" if kind == "warn" else "EEF3FA")
    p = c.paragraphs[0]; r = p.add_run(title); r.bold = True; r.font.color.rgb = NAVY; p.paragraph_format.space_after = Pt(2)
    p2 = c.add_paragraph(); add_runs(p2, text, 9.2); p2.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def picture(src, width):
    path = ROOT / src
    im = Image.open(path)
    max_h = Cm(15.5)
    w = width; h = int(w * im.height / im.width)
    if h > max_h:
        w = int(max_h * im.width / im.height)
    doc.add_picture(str(path), width=Emu(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True

# ---------- cover ----------
for _ in range(3): doc.add_paragraph()
doc.add_picture(str(ROOT.parent / "build" / "icon.png"), width=Cm(3))
para(meta["product"], size=40, color=NAVY, after=2).runs[0].bold = True
para(meta["tagline"], size=15, color=GREY, after=18)
para(f"{meta['release_designation']} · Version {meta['version']} · {meta['release_date']}", size=11, color=ACCENT, after=40)
for k, v in (("Company", meta["company"]), ("Maintainer", meta["founder"]), ("Website", meta["website"]), ("Contact", meta["email"]), ("Repository", meta["github"]), ("License", "Apache License 2.0"), ("Document revision", meta["revision"])):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{k}: "); r.bold = True; r.font.color.rgb = GREY; r.font.size = Pt(9)
    r2 = p.add_run(v); r2.font.size = Pt(9)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- contents ----------
cp = para("Contents", size=20, color=NAVY, after=10); cp.runs[0].bold = True
def hyperlink(par, anchor, text):
    h = OxmlElement("w:hyperlink"); h.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F2937"); rPr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "21"); rPr.append(sz)
    r.append(rPr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r)
    par._p.append(h)

for s in model["sections"]:
    row = doc.add_paragraph(); row.paragraph_format.space_after = Pt(3)
    nr = row.add_run(f"{s['num']}   "); nr.bold = True; nr.font.color.rgb = ACCENT; nr.font.size = Pt(10.5)
    hyperlink(row, f"s{s['num']}", s["title"])
para("Entries link to their chapters. To add page numbers in Word: References → Table of Contents.", size=8, color=GREY)
dp = para("Document information", size=13, color=NAVY, after=4); dp.runs[0].bold = True
table(["Field", "Value"], [["Product", f"{meta['product']} {meta['version']}"], ["Release", f"{meta['release_designation']}, {meta['release_date']}"], ["Document revision", meta["revision"]], ["Author", meta["author"]], ["License", meta["license"]], ["Distribution", meta["classification"]]], [30, 70])

# ---------- body ----------
body = doc.add_section(WD_SECTION.NEW_PAGE)
for s in (body,):
    hp = s.header.paragraphs[0]; hp.text = ""
    add_runs(hp, f"Cortextrace {meta['version']} · Product Release Document", 7.5, GREY)
    hp.add_run("\t\tPublic · Revision " + meta["revision"]).font.size = Pt(7.5)
    fp = s.footer.paragraphs[0]; fp.text = ""
    add_runs(fp, "© 2026 Manjunatha M, AI and ML consultants · Apache License 2.0\t\tPage ", 7.5, GREY)
    field(fp, "PAGE").font.size = Pt(7.5)
doc.sections[0].different_first_page_header_footer = True

first = True
for sec_model in model["sections"]:
    h = doc.add_heading(f"{sec_model['num']}  {sec_model['title']}", level=1)
    bid = str(100 + int(sec_model["num"]))
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), bid); bs.set(qn("w:name"), f"s{sec_model['num']}")
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), bid)
    h._p.insert(0, bs); h._p.append(be)
    if not first:
        h.paragraph_format.page_break_before = True
    first = False
    for b in sec_model["blocks"]:
        t = b["t"]
        if t == "h2": doc.add_heading(re.sub(r"[*`]", "", b["text"]), level=2)
        elif t == "h3": doc.add_heading(re.sub(r"[*`]", "", b["text"]), level=3)
        elif t == "p": para(b["text"])
        elif t == "lead": para(b["text"], size=12, color=RGBColor(0x37, 0x41, 0x51), after=10)
        elif t == "ul":
            for i in b["items"]: add_runs(doc.add_paragraph(style="List Bullet"), i)
        elif t == "ol":
            for n, i in enumerate(b["items"], 1):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
                r = p.add_run(f"{n}.\t"); r.bold = True; r.font.color.rgb = ACCENT
                add_runs(p, i)
        elif t == "table": table(b["head"], b["rows"], b.get("widths"))
        elif t == "note": note(b["title"], b["text"], b["kind"])
        elif t == "code":
            tt = doc.add_table(rows=1, cols=1); c = tt.rows[0].cells[0]; c.width = TEXT_WIDTH; shade(c, "0F1B2D")
            c.paragraphs[0].text = ""
            for li, line in enumerate(b["text"].split("\n")):
                pp = c.paragraphs[0] if li == 0 else c.add_paragraph()
                pp.paragraph_format.space_after = Pt(0)
                r = pp.add_run(line); r.font.name = MONO; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xE6, 0xED, 0xF6)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "feature":
            doc.add_heading(b["name"], level=3)
            ft = doc.add_table(rows=2, cols=2); ft.autofit = False; cell_borders(ft)
            for col, (k, v) in enumerate((("WHAT IT IS", b["what"]), ("WHY IT MATTERS", b["why"]))):
                kc = ft.rows[0].cells[col]; kc.width = int(TEXT_WIDTH / 2); shade(kc, "EEF3FA")
                kr = kc.paragraphs[0].add_run(k); kr.bold = True; kr.font.size = Pt(7.5); kr.font.color.rgb = ACCENT
                vc = ft.rows[1].cells[col]; vc.width = int(TEXT_WIDTH / 2); add_runs(vc.paragraphs[0], v, 9.2)
            hp = doc.add_paragraph(); hp.paragraph_format.space_before = Pt(6)
            hr = hp.add_run("HOW IT WORKS"); hr.bold = True; hr.font.size = Pt(7.5); hr.font.color.rgb = ACCENT
            for i in b["how"]: add_runs(doc.add_paragraph(style="List Bullet"), i, 9.5)
        elif t == "figure":
            picture(b["src"], int(TEXT_WIDTH * (0.92 if b["kind"] == "screenshot" else 1)))
            cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            m = re.match(r"(Figure \d+\.)\s*(.*)", b["caption"])
            r = cap.add_run(m.group(1) + " "); r.bold = True; r.font.color.rgb = NAVY
            add_runs(cap, m.group(2), 8.5, GREY)
            for a in b["legend"]:
                lp = doc.add_paragraph(); lp.paragraph_format.space_after = Pt(1); lp.paragraph_format.left_indent = Cm(0.7); lp.paragraph_format.first_line_indent = Cm(-0.7)
                nr = lp.add_run(f"{a['n']}\t"); nr.bold = True; nr.font.size = Pt(8.5); nr.font.color.rgb = ACCENT
                if a.get("label"):
                    lr = lp.add_run(a["label"] + " — "); lr.bold = True; lr.font.size = Pt(8.5)
                add_runs(lp, a["note"], 8.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

props = doc.core_properties
props.title = meta["title"]; props.author = meta["author"]; props.subject = f"{meta['product']} {meta['version']}"
props.keywords = "Cortextrace, AI agents, observability, security"; props.revision = 1
out = ROOT / f"Cortextrace_Product_Release_Document_v{meta['version']}.docx"
doc.save(out)
print("saved", out)
